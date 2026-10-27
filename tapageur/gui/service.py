"""Orchestration GUI : build items → pipeline → alertes → export."""
from __future__ import annotations

import json
import os
import sys
from collections import defaultdict
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from urllib.parse import urlparse

from tapageur.dedup import (apply_venue_defaults, canonicalize, convert_to_onliner,
                            dedup, group_events, load_mappings, load_venue_defaults)
from tapageur.gui import export_xlsx as _xlsx
from tapageur.gui.inputs import (TextInput, UrlInput, items_from_files,
                                 items_from_texts, items_from_urls)
from tapageur.normalize import normalize_items
from tapageur.schema import Event
from tapageur.validate import _source_ref_display, export_csv


def default_reference_year(today: date | None = None) -> str:
    """Année de référence pour le mode « tous les mois » : année courante,
    sauf en décembre (on prépare janvier de l'année suivante)."""
    today = today or date.today()
    return str(today.year + 1 if today.month == 12 else today.year)


@dataclass
class RunRequest:
    url_inputs: list[UrlInput]
    months: list[str]
    formats: list[str]            # sous-ensemble de {"xlsx", "csv", "json"}
    out_dir: Path
    tapeur_email: str
    month_mode: str = "specific"          # "specific" | "all"
    reference_year: str = ""              # mode "all" ; vide → default_reference_year()
    write_mode: str = "create"            # "create" | "append" | "merge"
    text_inputs: list[TextInput] = field(default_factory=list)
    file_inputs: list = field(default_factory=list)   # list[FileInput]


@dataclass
class MonthResult:
    month: str | None
    n_events: int
    files: list[Path] = field(default_factory=list)
    alerts: list[str] = field(default_factory=list)
    events: list = field(default_factory=list)   # list[Event] (détail pour l'UI)
    added: int = 0
    filled: int = 0
    unchanged: int = 0
    sources_skipped: int = 0


def build_alerts(events: list[Event], ref_cities: set[str], ref_venues: set[str]) -> list[str]:
    """Signale prix/performer manquants et lieux/villes hors listes de référence."""
    alerts: list[str] = []
    for ev in events:
        if ev.price == "?":
            alerts.append(f"Prix manquant : {ev.date_label} — {ev.acts[0].performer}")
        if "?" in ev.times:
            alerts.append(f"Horaire manquant : {ev.date_label} — {ev.acts[0].performer}")
        for act in ev.acts:
            if act.performer == "?":
                alerts.append(f"Artiste/performer inconnu : {ev.date_label} @ {ev.venue}")
        if ref_cities and ev.city and ev.city not in ref_cities:
            alerts.append(f"Ville hors liste de référence : {ev.city}")
        if ref_venues and ev.venue and ev.venue not in ref_venues:
            alerts.append(f"Lieu hors liste de référence : {ev.venue}")
    return alerts


def _reference_lists() -> tuple[set[str], set[str]]:
    """Lit communes / salles_lieux depuis le template embarqué."""
    import openpyxl
    wb = openpyxl.load_workbook(_xlsx.TEMPLATE_PATH, read_only=True)
    try:
        cities = {r[0] for r in wb["communes"].iter_rows(values_only=True)
                  if r and r[0] and not str(r[0]).startswith("=")}
        venues = {r[0] for r in wb["salles_lieux"].iter_rows(values_only=True)
                  if r and r[0] and not str(r[0]).startswith("=")}
    finally:
        wb.close()
    return cities, venues


def _mappings_path() -> Path:
    """config/mappings.yaml : embarqué sous sys._MEIPASS dans l'exe figé,
    relatif au CWD sinon (comportement CLI inchangé)."""
    if getattr(sys, "frozen", False):
        return Path(sys._MEIPASS) / "config" / "mappings.yaml"
    return Path("config/mappings.yaml")


def _venue_defaults_path() -> Path:
    if getattr(sys, "frozen", False):
        return Path(sys._MEIPASS) / "config" / "venue_defaults.yaml"
    return Path("config/venue_defaults.yaml")


def _pipeline(items, month: str, month_filter: bool, cache_out: Path,
              progress=None, reference_year: str | None = None) -> list[Event]:
    events = normalize_items(items, month, cache_out, model=None,
                             month_filter=month_filter, progress=progress,
                             reference_year=reference_year)
    events = canonicalize(events, load_mappings(_mappings_path()))
    events = apply_venue_defaults(events, load_venue_defaults(_venue_defaults_path()))
    events = dedup(events)
    events = group_events(events)
    onliner_names = [ev.festival for ev in events if ev.is_festival and ev.festival]
    events = convert_to_onliner(events, onliner_names)
    return events


def _input_source_display(kind: str, ref: str) -> str:
    """Référence SOURCE telle qu'écrite dans le xlsx (miroir validate)."""
    return _source_ref_display(kind, ref)


def _filter_present_sources(req: "RunRequest", target: Path, emit) -> int:
    """En mode incrémental : retire les file/url inputs dont la SOURCE est déjà
    dans le fichier cible. Retourne le nombre de sources ignorées. Textes gardés
    (pas d'identifiant stable)."""
    present = _xlsx.existing_sources(target)
    if not present:
        return 0
    present_norm = {p.strip() for p in present}
    # Fichiers : comparaison insensible à la casse (noms de fichiers Windows).
    present_ci = {p.casefold() for p in present_norm}
    # URLs à suivre : l'URL d'agenda saisie n'est PAS stockée (la colonne SOURCE
    # garde les pages détail). On rapproche donc par HÔTE : si l'hôte de l'URL
    # saisie apparaît déjà dans les sources, on considère la source présente.
    # (Interim ; le fix exact = colonne « id source config ».)
    present_hosts = {h for h in (urlparse(s).netloc for s in present_norm) if h}
    skipped = 0
    kept_files = []
    for fi in req.file_inputs:
        disp = _input_source_display("image", fi.path)  # basename, suffisant pour fichiers/pdf
        base = Path(fi.path).name
        if base.casefold() in present_ci or disp.casefold() in present_ci:
            skipped += 1
        else:
            kept_files.append(fi)
    kept_urls = []
    for u in req.url_inputs:
        host = urlparse(u.url).netloc
        follow = getattr(u, "follow_links", False)
        if u.url in present_norm or (follow and host and host in present_hosts):
            skipped += 1
        else:
            kept_urls.append(u)
    req.file_inputs = kept_files
    req.url_inputs = kept_urls
    if skipped:
        emit(f"Incrémental : {skipped} source(s) déjà présente(s) ignorée(s).")
    return skipped


def _text_id(text: str) -> str:
    """Id stable d'un texte collé = hash de son contenu. Permet au manifeste de
    reconnaître un texte déjà traité (vs l'ancien « texte #N » positionnel qui,
    toujours retraité, dupliquait ses events à chaque passe « Ajouter »)."""
    import hashlib
    return "texte:" + hashlib.sha1(text.strip().encode("utf-8")).hexdigest()[:10]


def _input_source_ids(req: "RunRequest") -> list:
    """(id, type) des sources saisies, pour le manifeste. id = URL (web),
    basename (fichier), « texte:<hash> » (texte collé)."""
    ids: list = []
    for fi in req.file_inputs:
        ids.append((Path(fi.path).name, "Fichier"))
    for u in req.url_inputs:
        ids.append((u.url, "URL"))
    for t in req.text_inputs:
        txt = (getattr(t, "text", "") or "").strip()
        if txt:
            ids.append((_text_id(txt), "Texte"))
    return ids


def _filter_by_manifest(req: "RunRequest", target: Path, emit) -> int:
    """Skip via manifeste : retire les file/url/text inputs déjà traités.
    Fichiers insensibles à la casse ; urls exact ; textes par hash de contenu
    (un texte identique déjà traité n'est pas ré-ajouté → pas de doublon)."""
    manifest = _xlsx.read_processed_sources(target)
    if not manifest:
        return 0
    manifest_ci = {m.casefold() for m in manifest}
    skipped = 0
    kept_files = []
    for fi in req.file_inputs:
        if Path(fi.path).name.casefold() in manifest_ci:
            skipped += 1
        else:
            kept_files.append(fi)
    kept_urls = []
    for u in req.url_inputs:
        if u.url in manifest:
            skipped += 1
        else:
            kept_urls.append(u)
    kept_texts = []
    for t in req.text_inputs:
        txt = (getattr(t, "text", "") or "").strip()
        if txt and _text_id(txt) in manifest:
            skipped += 1
        else:
            kept_texts.append(t)
    req.file_inputs = kept_files
    req.url_inputs = kept_urls
    req.text_inputs = kept_texts
    if skipped:
        emit(f"Incrémental : {skipped} source(s) déjà traitée(s) ignorée(s).")
    return skipped


def _export(events: list[Event], out_dir: Path, label: str, formats: list[str],
            write_mode: str = "create") -> tuple[list[Path], "_xlsx.MergeReport | None"]:
    out_dir.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    report = None
    if "json" in formats:
        p = out_dir / f"tapage_{label}.json"
        p.write_text(json.dumps([e.model_dump() for e in events], indent=2, ensure_ascii=False),
                     encoding="utf-8")
        written.append(p)
    if "csv" in formats:
        p = out_dir / f"tapage_{label}.csv"
        export_csv(events, p)
        written.append(p)
    if "xlsx" in formats:
        p = out_dir / f"tapage_{label}.xlsx"
        if write_mode == "merge":
            report = _xlsx._merge_xlsx(events, p)
        elif write_mode == "append":
            _xlsx.export_xlsx(events, p, append=True)
        else:
            _xlsx.export_xlsx(events, p)
        written.append(p)
    return written, report


def run(req: RunRequest, progress: Callable[[str], None] | None = None) -> list[MonthResult]:
    def _emit(msg):
        progress and progress(msg)

    os.environ["TAPEUR_EMAIL"] = req.tapeur_email
    ref_cities, ref_venues = _reference_lists()
    all_input_ids = _input_source_ids(req)

    def _apply_report(mr: MonthResult, report) -> None:
        if report is None:
            return
        mr.added, mr.filled, mr.unchanged = report.added, report.filled, report.unchanged
        _emit(f"Fusion : {report.added} ajoutés · {report.filled} complétés · "
              f"{report.unchanged} inchangés.")
        for k in report.added_keys:
            _emit(f"  + AJOUTÉ  {k[0]} · {k[1]} · {k[2]}")
        for k, labels in report.filled_detail:
            _emit(f"  ~ COMPLÉTÉ {k[0]} · {k[1]} · {k[2]} → {', '.join(labels)}")

    # Skip incrémental : uniquement en mode spécifique mono-mois + write_mode append.
    sources_skipped = 0
    if req.month_mode != "all" and req.write_mode == "append" and len(req.months) == 1:
        target = req.out_dir / f"tapage_{req.months[0]}.xlsx"
        if _xlsx.has_manifest(target):
            sources_skipped = _filter_by_manifest(req, target, _emit)
        else:
            sources_skipped = _filter_present_sources(req, target, _emit)  # repli

    _emit("Collecte des sources…")
    url_month = "" if req.month_mode == "all" else (req.months[0] if req.months else "")
    items = list(items_from_files(req.file_inputs, progress=progress))
    items += list(items_from_urls(req.url_inputs, url_month, progress=progress))
    items += list(items_from_texts(req.text_inputs, progress=progress))

    results: list[MonthResult] = []

    if req.month_mode == "all":
        ref_year = req.reference_year or default_reference_year()
        key = f"tous-{ref_year}"
        cache_out = Path("cache") / key / "gui_normalized.json"
        _emit("tous les mois : extraction…")
        events = _pipeline(items, key, False, cache_out, progress=progress,
                           reference_year=ref_year)
        by_month: dict[str, list[Event]] = defaultdict(list)
        for ev in events:
            by_month[ev.date_iso[:7]].append(ev)
        for month in sorted(by_month):
            evs = by_month[month]
            _emit(f"{month} : {len(evs)} events, export…")
            files, report = _export(evs, req.out_dir, month, req.formats,
                                     write_mode=req.write_mode)
            if "xlsx" in req.formats:
                _xlsx.record_processed_sources(req.out_dir / f"tapage_{month}.xlsx",
                                               all_input_ids)
            alerts = build_alerts(evs, ref_cities, ref_venues)
            if req.write_mode in ("append", "merge") and "xlsx" not in req.formats:
                alerts.append(f"Mode '{req.write_mode}' ignoré : xlsx non sélectionné.")
            mr = MonthResult(month=month, n_events=len(evs),
                             files=files, alerts=alerts, events=evs)
            _apply_report(mr, report)
            results.append(mr)
        if not by_month:
            results.append(MonthResult(month=None, n_events=0, files=[],
                                       alerts=["Aucun event extrait."]))
    else:
        for month in req.months:
            cache_out = Path("cache") / month / "gui_normalized.json"
            _emit(f"{month} : extraction…")
            events = _pipeline(items, month, True, cache_out, progress=progress)
            _emit(f"{month} : {len(events)} events, export…")
            files, report = _export(events, req.out_dir, month, req.formats,
                                     write_mode=req.write_mode)
            if "xlsx" in req.formats:
                _xlsx.record_processed_sources(req.out_dir / f"tapage_{month}.xlsx",
                                               all_input_ids)
            alerts = build_alerts(events, ref_cities, ref_venues)
            if req.write_mode in ("append", "merge") and "xlsx" not in req.formats:
                alerts.append(f"Mode '{req.write_mode}' ignoré : xlsx non sélectionné.")
            mr = MonthResult(month=month, n_events=len(events),
                             files=files, alerts=alerts, events=events,
                             sources_skipped=sources_skipped)
            _apply_report(mr, report)
            results.append(mr)

    _emit("Terminé.")
    return results
