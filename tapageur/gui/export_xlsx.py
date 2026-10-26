"""Export xlsx : remplit une copie du template (feuille `lebiduleur`)."""
from __future__ import annotations

import os
import sys
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from tapageur.schema import Event
from tapageur.validate import event_to_row

# Idem app._WEB : résoudre sous sys._MEIPASS quand figé (exe PyInstaller).
_GUI_DIR = (Path(sys._MEIPASS) / "tapageur" / "gui"
            if getattr(sys, "frozen", False)
            else Path(__file__).resolve().parent)
TEMPLATE_PATH = _GUI_DIR / "assets" / "template.xlsx"
DATA_SHEET = "lebiduleur"
_N_FIXED_COLS = 12  # colonnes A-L (TYPE SOURCE, SOURCE, puis INFO TAPEUR … PRIX)


def export_xlsx(events: list[Event], out_path: Path, template_path: Path = TEMPLATE_PATH,
                append: bool = False) -> None:
    """Copie le template et écrit les events dans la feuille `lebiduleur`.

    Blocs ACT dynamiques = max(len(acts)). Les en-têtes des blocs au-delà de 4
    (ceux présents dans le template) sont ajoutés. Les feuilles de référence
    sont laissées intactes.

    append=False (ou fichier absent) : part d'une copie du template (création).
    append=True + fichier existant : ajoute les lignes à la suite, en préservant
    les lignes et feuilles existantes."""
    if append and out_path.exists():
        _append_xlsx(events, out_path)
        return

    wb = openpyxl.load_workbook(template_path)
    ws = wb[DATA_SHEET]

    # Purge d'éventuelles lignes de données résiduelles (garde la ligne 1).
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    n_acts = max((len(ev.acts) for ev in events), default=1)

    # En-têtes des blocs ACT supplémentaires (au-delà des 4 du template).
    _write_act_headers(ws, 5, n_acts)

    for r, ev in enumerate(events, start=2):
        for c, value in enumerate(event_to_row(ev, n_acts=n_acts), start=1):
            ws.cell(row=r, column=c, value=value)

    _save_workbook(wb, out_path)


def _save_workbook(wb, out_path: Path) -> None:
    """Sauvegarde atomique : écrit dans un fichier temporaire puis le substitue.
    Si la cible est verrouillée (fichier ouvert dans Excel…), lève une
    PermissionError avec un message clair, sans toucher au fichier existant."""
    out_path.parent.mkdir(parents=True, exist_ok=True)
    tmp = out_path.with_name(out_path.name + ".tmp")
    wb.save(tmp)
    try:
        os.replace(tmp, out_path)
    except PermissionError as e:
        try:
            tmp.unlink()
        except OSError:
            pass
        raise PermissionError(
            f"Impossible d'écrire {out_path.name} : le fichier est peut-être ouvert "
            f"dans Excel. Ferme-le puis relance."
        ) from e


def _write_act_headers(ws, first: int, last: int) -> None:
    """Écrit les en-têtes des blocs ACT `first..last` (inclus) sur la ligne 1."""
    for i in range(first, last + 1):
        base = _N_FIXED_COLS + (i - 1) * 4
        ws.cell(row=1, column=base + 1, value=f"GENRE {i}")
        ws.cell(row=1, column=base + 2, value=f"NOM SPECTACLE {i}")
        ws.cell(row=1, column=base + 3, value=f"ARTISTE {i}")
        ws.cell(row=1, column=base + 4, value=f"STYLE {i}")


def _append_xlsx(events: list[Event], out_path: Path) -> None:
    wb = openpyxl.load_workbook(out_path)
    if DATA_SHEET not in wb.sheetnames:
        raise ValueError(
            f"{out_path} n'a pas de feuille '{DATA_SHEET}' : append impossible."
        )
    ws = wb[DATA_SHEET]
    existing_blocks = max(1, (ws.max_column - _N_FIXED_COLS) // 4)
    n_acts = max(existing_blocks, max((len(ev.acts) for ev in events), default=1))
    if n_acts > existing_blocks:
        _write_act_headers(ws, existing_blocks + 1, n_acts)
    start = ws.max_row + 1
    for r, ev in enumerate(events, start=start):
        for c, value in enumerate(event_to_row(ev, n_acts=n_acts), start=1):
            ws.cell(row=r, column=c, value=value)
    _save_workbook(wb, out_path)


# --- Primitives merge/dédoublonnage (clé de match sur colonnes visibles) ---
# (unicodedata + event_to_row sont importés en tête du module)

_COL_TYPE_SOURCE = 1
_COL_SOURCE = 2
_COL_FESTOCHE = 6
_COL_STYLE_FESTOCHE = 7
_COL_DATE = 8
_COL_HEURE = 9
_COL_LIEU = 10
_COL_VILLE = 11
_COL_PRIX = 12
_COL_ARTISTE1 = _N_FIXED_COLS + 3  # 1er bloc ACT : GENRE=13,NOM=14,ARTISTE=15,STYLE=16


def _norm(v) -> str:
    s = "" if v is None else str(v)
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    return s.strip().lower()


def _is_blank(v) -> bool:
    return v is None or str(v).strip() in ("", "?")


def _cell(cells: list, col: int):
    return cells[col - 1] if 0 <= col - 1 < len(cells) else ""


def _key_from_cells(cells: list) -> tuple:
    return (_norm(_cell(cells, _COL_DATE)),
            _norm(_cell(cells, _COL_LIEU)),
            _norm(_cell(cells, _COL_ARTISTE1)))


def _event_cells(ev, n_acts: int) -> list:
    return event_to_row(ev, n_acts=n_acts)


# --- Fill-gaps : plan de complétion d'une ligne existante ---

_KEY_COLS = {_COL_DATE, _COL_LIEU, _COL_ARTISTE1}
_META_COLS = {_COL_TYPE_SOURCE, _COL_SOURCE, 3, 4, 5}  # TYPE SOURCE, SOURCE, INFO TAPEUR, DATE TAPAGE, DATE EMISSION


def _label_for(col: int) -> str:
    fixed = {
        _COL_FESTOCHE: "FESTOCHE", _COL_STYLE_FESTOCHE: "STYLE FESTOCHE",
        _COL_HEURE: "HEURE", _COL_VILLE: "VILLE", _COL_PRIX: "PRIX",
    }
    if col in fixed:
        return fixed[col]
    # bloc ACT : i = (col-1 - _N_FIXED_COLS)//4 + 1 ; sous-champ selon offset
    offset = (col - 1 - _N_FIXED_COLS) % 4
    i = (col - 1 - _N_FIXED_COLS) // 4 + 1
    sub = {0: "GENRE", 1: "NOM", 2: "ARTISTE", 3: "STYLE"}[offset]
    return f"{sub} {i}"


def _plan_fill(existing: list, incoming: list, n_cols: int):
    fills: dict[int, object] = {}
    labels: list[str] = []
    for col in range(1, n_cols + 1):
        if col in _KEY_COLS or col in _META_COLS:
            continue
        ex = _cell(existing, col)
        inc = _cell(incoming, col)
        if _is_blank(ex) and not _is_blank(inc):
            fills[col] = inc
            labels.append(_label_for(col))
    return fills, labels


def _merge_source_cell(existing_val, incoming_val) -> str | None:
    inc = ("" if incoming_val is None else str(incoming_val)).strip()
    if not inc:
        return None
    parts = [p.strip() for p in str(existing_val or "").split(";") if p.strip()]
    if any(_norm(p) == _norm(inc) for p in parts):
        return None
    parts.append(inc)
    return " ; ".join(parts)


# --- Fusion in-place + compte-rendu ---


@dataclass
class MergeReport:
    added: int = 0
    filled: int = 0
    unchanged: int = 0
    added_keys: list = field(default_factory=list)       # list[tuple]
    filled_detail: list = field(default_factory=list)     # list[(key, [labels])]


def _merge_xlsx(events: list, out_path: Path) -> MergeReport:
    """Fusionne `events` dans un xlsx existant : complète les trous des lignes
    qui matchent (DATE+LIEU+ARTISTE1), ajoute les non-matchés. In-place."""
    if not out_path.exists():
        export_xlsx(events, out_path)  # création
        return MergeReport(added=len(events),
                           added_keys=[_key_from_cells(_event_cells(e, max(1, len(e.acts))))
                                       for e in events])

    wb = openpyxl.load_workbook(out_path)
    if DATA_SHEET not in wb.sheetnames:
        raise ValueError(f"{out_path} n'a pas de feuille '{DATA_SHEET}' : fusion impossible.")
    ws = wb[DATA_SHEET]

    existing_blocks = max(1, (ws.max_column - _N_FIXED_COLS) // 4)
    target_blocks = max(existing_blocks, max((len(e.acts) for e in events), default=1))
    if target_blocks > existing_blocks:
        _write_act_headers(ws, existing_blocks + 1, target_blocks)
    n_cols = _N_FIXED_COLS + target_blocks * 4

    # Index des lignes existantes par clé (1re occurrence).
    row_cells: dict[int, list] = {}
    key_to_row: dict[tuple, int] = {}
    for r in range(2, ws.max_row + 1):
        cells = [ws.cell(row=r, column=c).value for c in range(1, n_cols + 1)]
        k = _key_from_cells(cells)
        if k == ("", "", ""):
            continue  # ligne vide
        row_cells[r] = cells
        key_to_row.setdefault(k, r)

    report = MergeReport()
    to_append: list = []
    append_key_idx: dict[tuple, int] = {}   # clé → index dans to_append (dédup entrants)
    matched_existing: set = set()           # lignes existantes touchées par ≥1 entrant
    touched: dict[int, list] = {}           # ligne existante → libellés remplis (cumulés)

    def _compute_changes(target: list, inc: list):
        """{col: val} à appliquer (trous de contenu + SOURCE/TYPE accolées) et libellés."""
        fills, labels = _plan_fill(target, inc, n_cols)
        changes = dict(fills)
        ns = _merge_source_cell(target[_COL_SOURCE - 1], inc[_COL_SOURCE - 1])
        if ns is not None:
            changes[_COL_SOURCE] = ns
        nt = _merge_source_cell(target[_COL_TYPE_SOURCE - 1], inc[_COL_TYPE_SOURCE - 1])
        if nt is not None:
            changes[_COL_TYPE_SOURCE] = nt
        return changes, labels

    for ev in events:
        inc = _event_cells(ev, target_blocks)
        k = _key_from_cells(inc)
        r = key_to_row.get(k)
        if r is not None:
            # Match sur une ligne existante. Le snapshot `existing` est rafraîchi
            # au fil des entrants → un 2e entrant de même clé ne clobbe pas et ne
            # perd pas la source du 1er (il comble seulement de nouveaux trous).
            matched_existing.add(r)
            existing = row_cells[r]
            changes, labels = _compute_changes(existing, inc)
            for col, val in changes.items():
                ws.cell(row=r, column=col, value=val)
                existing[col - 1] = val
            if changes:
                acc = touched.setdefault(r, [])
                for lab in labels:
                    if lab not in acc:
                        acc.append(lab)
            continue
        j = append_key_idx.get(k)
        if j is not None:
            # 2e entrant de même clé absent du fichier : fusionner dans la ligne
            # en attente plutôt que d'ajouter un doublon.
            pending = to_append[j]
            changes, _ = _compute_changes(pending, inc)
            for col, val in changes.items():
                pending[col - 1] = val
            continue
        append_key_idx[k] = len(to_append)
        to_append.append(inc)
        report.added += 1
        report.added_keys.append(k)

    # Append des non-matchés (dédupliqués entre entrants) à la suite.
    start = ws.max_row + 1
    for i, inc in enumerate(to_append):
        for c, value in enumerate(inc, start=1):
            ws.cell(row=start + i, column=c, value=value)

    # Compteurs : une ligne existante est « complétée » si ≥1 changement, sinon
    # « inchangée » (comptée une seule fois quel que soit le nombre d'entrants).
    report.filled = len(touched)
    report.unchanged = len(matched_existing - set(touched))
    row_key = {r: k for k, r in key_to_row.items()}
    report.filled_detail = [(row_key.get(r, ("", "", "")), labs) for r, labs in touched.items()]

    _save_workbook(wb, out_path)
    return report


# --- Lecture des sources présentes (skip incrémental) ---


def existing_sources(out_path: Path) -> set:
    """Ensemble des sources déjà présentes (colonne SOURCE), pour le skip
    incrémental. Suffixe ' (p.N)' des PDF retiré. Vide si fichier absent."""
    if not out_path.exists():
        return set()
    try:
        wb = openpyxl.load_workbook(out_path, read_only=True, data_only=True)
    except Exception:
        return set()
    try:
        if DATA_SHEET not in wb.sheetnames:
            return set()
        ws = wb[DATA_SHEET]
        out: set = set()
        for r in ws.iter_rows(min_row=2, min_col=_COL_SOURCE, max_col=_COL_SOURCE,
                              values_only=True):
            val = r[0]
            if not val:
                continue
            for part in str(val).split(";"):
                s = part.strip()
                if s.endswith(")") and " (p." in s:
                    s = s[:s.rindex(" (p.")]
                if s:
                    out.add(s)
        return out
    finally:
        wb.close()
