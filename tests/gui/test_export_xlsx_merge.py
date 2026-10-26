# tests/gui/test_export_xlsx_merge.py
from tapageur.gui import export_xlsx as X
from tapageur.schema import Event, Act


def _ev(perf="Groupe X", venue="Le Lézard", date_label="Mardi 3",
        heure="20h00", price="?", festival=None, is_festival=False):
    return Event(
        tapeur_email="t@b.fr", festival=festival, date_label=date_label,
        date_iso="2026-11-03", times=[heure], venue=venue, city="Le Mans",
        price=price, is_festival=is_festival,
        acts=[Act(genre="c", title=None, performer=perf, style=None)],
        source_kind="web", source_ref="https://a.fr/x",
    )


def test_norm_strips_accents_case_space():
    assert X._norm("  Le Lézard ") == X._norm("le lezard")


def test_is_blank():
    assert X._is_blank(None) and X._is_blank("") and X._is_blank("?") and X._is_blank(" ? ")
    assert not X._is_blank("20h00")


def test_key_from_cells_matches_event_cells():
    ev = _ev()
    cells = X._event_cells(ev, n_acts=1)
    # DATE col8, LIEU col10, ARTISTE1 col15
    assert X._key_from_cells(cells) == ("mardi 3", "le lezard", "groupe x")


def test_key_ignores_heure_price():
    a = X._key_from_cells(X._event_cells(_ev(heure="20h00", price="10€"), 1))
    b = X._key_from_cells(X._event_cells(_ev(heure="21h30", price="?"), 1))
    assert a == b


def test_plan_fill_fills_blanks_only():
    ev_existing = _ev(heure="20h00", price="?", festival=None)
    ev_incoming = _ev(heure="21h00", price="10€", festival="Soirée Beaujolais")
    existing = X._event_cells(ev_existing, 1)
    incoming = X._event_cells(ev_incoming, 1)
    fills, labels = X._plan_fill(existing, incoming, n_cols=len(existing))
    # PRIX (col12) vide → rempli ; FESTOCHE (col6) vide → rempli
    assert fills[X._COL_PRIX] == "10€"
    assert fills[X._COL_FESTOCHE] == "Soirée Beaujolais"
    # HEURE déjà remplie (20h00) → NON touchée
    assert X._COL_HEURE not in fills
    assert "PRIX" in labels and "FESTOCHE" in labels


def test_plan_fill_never_overwrites():
    existing = X._event_cells(_ev(heure="20h00", price="10€"), 1)
    incoming = X._event_cells(_ev(heure="21h00", price="15€"), 1)
    fills, labels = X._plan_fill(existing, incoming, n_cols=len(existing))
    assert fills == {} and labels == []


def test_merge_source_cell_appends_without_dup():
    assert X._merge_source_cell("prog.pdf", "https://a.fr/x") == "prog.pdf ; https://a.fr/x"
    assert X._merge_source_cell("https://a.fr/x", "https://a.fr/x") is None  # déjà présent


from pathlib import Path


def test_merge_fills_match_and_appends_new(tmp_path: Path):
    out = tmp_path / "tapage_2026-11.xlsx"
    # Fichier initial : 1 event, HEURE connue, PRIX "?"
    base = _ev(perf="Groupe X", heure="20h00", price="?")
    X.export_xlsx([base], out)  # création
    # Fusion : même event (complète PRIX) + 1 nouvel event
    same = _ev(perf="Groupe X", heure="21h00", price="12€", festival="Soirée Beaujolais")
    new = _ev(perf="Autre Groupe", date_label="Vendredi 6", heure="20h30", price="gratuit")
    report = X._merge_xlsx([same, new], out)

    assert report.added == 1 and report.filled == 1 and report.unchanged == 0

    import openpyxl
    wb = openpyxl.load_workbook(out)
    ws = wb[X.DATA_SHEET]
    rows = {(r[7], r[9], r[14]): r for r in ws.iter_rows(min_row=2, values_only=True)}
    # ligne complétée : PRIX (col12→idx11) rempli, HEURE (idx8) préservée à 20h00
    g = rows[("Mardi 3", "Le Lézard", "Groupe X")]
    assert g[8] == "20h00"      # HEURE préservée
    assert g[11] == "12€"       # PRIX complété
    assert g[5] == "Soirée Beaujolais"  # FESTOCHE complété
    # nouvelle ligne ajoutée
    assert ("Vendredi 6", "Le Lézard", "Autre Groupe") in rows


def test_merge_absent_file_behaves_like_create(tmp_path: Path):
    out = tmp_path / "tapage_2026-11.xlsx"
    report = X._merge_xlsx([_ev()], out)
    assert out.exists() and report.added == 1 and report.filled == 0


def test_merge_appends_source(tmp_path: Path):
    out = tmp_path / "f.xlsx"
    base = _ev(perf="Groupe X", price="?")
    base.source_ref = "https://natif.fr/x"
    X.export_xlsx([base], out)
    same = _ev(perf="Groupe X", price="10€")
    same.source_ref = "https://clappin.fr/agenda-musique/"
    X._merge_xlsx([same], out)
    import openpyxl
    ws = openpyxl.load_workbook(out)[X.DATA_SHEET]
    src = [r[1] for r in ws.iter_rows(min_row=2, values_only=True)][0]
    assert "natif.fr" in src and "clappin.fr" in src


def test_existing_sources_reads_source_column(tmp_path: Path):
    out = tmp_path / "f.xlsx"
    e1 = _ev(); e1.source_ref = "https://clappin.fr/agenda-musique/"
    e2 = _ev(perf="Y", date_label="Jeudi 5")
    e2.source_kind = "image"; e2.source_ref = r"C:\x\flyer.png"
    X.export_xlsx([e1, e2], out)
    srcs = X.existing_sources(out)
    assert "https://clappin.fr/agenda-musique/" in srcs
    assert "flyer.png" in srcs  # image → basename


def test_existing_sources_absent_file():
    assert X.existing_sources(Path("nope-xyz.xlsx")) == set()


# --- Collisions de clé entre events ENTRANTS (Finding revue) ---

def test_merge_two_incoming_same_key_no_match_appended_once(tmp_path: Path):
    """Deux entrants même clé (DATE+LIEU+ARTISTE) absents du fichier → UNE seule
    ligne ajoutée (fusionnée), pas un doublon ; les trous sont complétés."""
    out = tmp_path / "f.xlsx"
    X.export_xlsx([_ev(perf="Base", date_label="Lundi 1", price="12€")], out)
    a = _ev(perf="Dup", date_label="Mardi 3", heure="20h00", price="?")
    a.source_ref = "https://a.fr/x"
    b = _ev(perf="Dup", date_label="Mardi 3", heure="21h00", price="10€")
    b.source_ref = "https://b.fr/y"
    report = X._merge_xlsx([a, b], out)
    assert report.added == 1
    import openpyxl
    ws = openpyxl.load_workbook(out)[X.DATA_SHEET]
    dup = [r for r in ws.iter_rows(min_row=2, values_only=True)
           if r[14] == "Dup"]
    assert len(dup) == 1                      # un seul doublon, pas deux
    assert dup[0][11] == "10€"                # trou PRIX comblé par le 2e
    assert "a.fr" in dup[0][1] and "b.fr" in dup[0][1]  # sources accolées


def test_merge_two_incoming_same_key_match_existing_no_clobber(tmp_path: Path):
    """Deux entrants même clé matchant une ligne existante : pas d'écrasement,
    les deux sources accolées, filled compté UNE fois."""
    out = tmp_path / "f.xlsx"
    base = _ev(perf="X", heure="20h00", price="?")
    base.source_ref = "https://natif.fr/x"
    X.export_xlsx([base], out)
    a = _ev(perf="X", heure="20h00", price="10€")
    a.source_ref = "https://a.fr/x"
    b = _ev(perf="X", heure="20h00", price="20€")
    b.source_ref = "https://b.fr/y"
    report = X._merge_xlsx([a, b], out)
    assert report.filled == 1 and report.added == 0
    import openpyxl
    ws = openpyxl.load_workbook(out)[X.DATA_SHEET]
    row = [r for r in ws.iter_rows(min_row=2, values_only=True)][0]
    assert row[8] == "20h00"                  # HEURE préservée
    assert row[11] == "10€"                   # 1er comble le trou, 2e n'écrase PAS
    assert all(s in row[1] for s in ("natif.fr", "a.fr", "b.fr"))  # 3 sources
