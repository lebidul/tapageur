import openpyxl
from tapageur.gui import export_xlsx
from tapageur.schema import Event, Act


def _event(perf="Groupe A", n_acts=1, festival=False):
    acts = [Act(genre="c", title=None, performer=f"{perf}{i}", style=None) for i in range(n_acts)]
    return Event(
        tapeur_email="t@b.fr", festival=("Fest" if festival else None),
        date_label="Vendredi 2", date_iso="2026-10-02", times=["20h00"],
        venue="Le Lieu", city="Le Mans", price="10€",
        is_festival=festival, acts=acts, source_kind="web", source_ref="https://x.fr",
    )


def test_export_writes_data_and_keeps_reference_sheets(tmp_path):
    out = tmp_path / "tapage.xlsx"
    export_xlsx.export_xlsx([_event()], out)
    wb = openpyxl.load_workbook(out)
    assert set(["lebiduleur", "communes", "salles_lieux", "date"]).issubset(set(wb.sheetnames))
    ws = wb["lebiduleur"]
    # En-têtes : provenance en A/B, INFO TAPEUR décalé en C
    assert ws.cell(row=1, column=1).value == "TYPE SOURCE"
    assert ws.cell(row=1, column=2).value == "SOURCE"
    assert ws.cell(row=1, column=3).value == "INFO TAPEUR"
    # Ligne 2 = 1er event
    assert ws.cell(row=2, column=1).value == "URL"            # TYPE SOURCE (source_kind=web)
    assert ws.cell(row=2, column=2).value == "https://x.fr"   # SOURCE
    assert ws.cell(row=2, column=3).value == "t@b.fr"         # INFO TAPEUR
    assert ws.cell(row=2, column=10).value == "Le Lieu"       # LIEU (col J, +2)
    assert ws.cell(row=2, column=11).value == "Le Mans"       # VILLE (col K, +2)


def test_export_dynamic_act_blocks(tmp_path):
    out = tmp_path / "tapage.xlsx"
    export_xlsx.export_xlsx([_event(n_acts=5)], out)
    ws = openpyxl.load_workbook(out)["lebiduleur"]
    # 12 colonnes fixes + 5 blocs * 4 = 32 colonnes ; en-tête bloc 5 ajouté
    assert ws.cell(row=1, column=13 + 4 * 4).value == "GENRE 5"
    assert ws.cell(row=2, column=13 + 4 * 4).value == "c"     # genre du 5e act


def test_append_preserves_and_adds(tmp_path):
    out = tmp_path / "tapage.xlsx"
    export_xlsx.export_xlsx([_event(perf="Alpha")], out)                 # création
    export_xlsx.export_xlsx([_event(perf="Beta")], out, append=True)     # append
    ws = openpyxl.load_workbook(out)["lebiduleur"]
    assert ws.cell(row=2, column=15).value == "Alpha0"   # ARTISTE 1, 1re ligne inchangée
    assert ws.cell(row=3, column=15).value == "Beta0"    # ARTISTE 1, ligne ajoutée
    assert ws.max_row == 3


def test_append_widens_act_blocks(tmp_path):
    out = tmp_path / "tapage.xlsx"
    export_xlsx.export_xlsx([_event(n_acts=1)], out)                     # 4 blocs (template)
    export_xlsx.export_xlsx([_event(n_acts=5, perf="X")], out, append=True)
    ws = openpyxl.load_workbook(out)["lebiduleur"]
    assert ws.cell(row=1, column=13 + 4 * 4).value == "GENRE 5"         # en-tête bloc 5 ajouté
    assert ws.cell(row=3, column=13 + 4 * 4).value == "c"               # genre 5e act, ligne ajoutée


def test_append_without_lebiduleur_raises(tmp_path):
    import openpyxl as _ox
    foreign = tmp_path / "foreign.xlsx"
    wb = _ox.Workbook()
    wb.active.title = "autre"
    wb.save(foreign)
    import pytest
    with pytest.raises(ValueError):
        export_xlsx.export_xlsx([_event()], foreign, append=True)


def test_append_on_missing_file_creates(tmp_path):
    out = tmp_path / "nouveau.xlsx"
    export_xlsx.export_xlsx([_event(perf="Solo")], out, append=True)     # fichier absent
    ws = openpyxl.load_workbook(out)["lebiduleur"]
    assert ws.cell(row=2, column=15).value == "Solo0"
    assert ws.max_row == 2


def test_save_locked_file_raises_clear_error(tmp_path, monkeypatch):
    import os as _os
    import pytest
    out = tmp_path / "tapage.xlsx"
    export_xlsx.export_xlsx([_event(perf="Un")], out)          # crée le fichier
    # Simule un fichier verrouillé (ouvert dans Excel) : os.replace échoue.
    monkeypatch.setattr(export_xlsx.os, "replace",
                        lambda *a, **k: (_ for _ in ()).throw(PermissionError("locked")))
    with pytest.raises(PermissionError) as exc:
        export_xlsx.export_xlsx([_event(perf="Deux")], out, append=True)
    assert "Excel" in str(exc.value)
    # Le fichier existant est préservé et aucun .tmp ne traîne.
    assert out.exists()
    assert not (tmp_path / "tapage.xlsx.tmp").exists()


def test_save_atomic_no_tmp_left(tmp_path):
    out = tmp_path / "tapage.xlsx"
    export_xlsx.export_xlsx([_event(perf="Solo")], out)
    assert out.exists()
    assert not (tmp_path / "tapage.xlsx.tmp").exists()
